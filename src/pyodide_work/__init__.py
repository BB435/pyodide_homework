
from io import BytesIO

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def process_image(image_bytes: bytes) -> bytes:
    """画像を解析し、結果をExcelのバイト列として返す。"""

    # ブラウザから渡された画像データをデコード
    buffer = np.frombuffer(image_bytes, dtype=np.uint8)
    image = cv2.imdecode(buffer, cv2.IMREAD_COLOR)

    if image is None:
        raise ValueError("画像を読み込めませんでした")

    # グレースケール変換
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)

    # 二値化（大津の方法）
    _, binary = cv2.threshold(
        gray,
        0,
        255,
        cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU,
    )

    # 輪郭検出
    contours, _ = cv2.findContours(
        binary,
        cv2.RETR_EXTERNAL,
        cv2.CHAIN_APPROX_SIMPLE,
    )

    # Excel作成
    wb = Workbook()
    ws = wb.active
    ws.title = "検出結果"

    headers = ["No", "X", "Y", "幅", "高さ", "面積"]
    ws.append(headers)

    for contour in contours:
        area = cv2.contourArea(contour)

        if area < 100:
            continue

        x, y, w, h = cv2.boundingRect(contour)

        ws.append([
            ws.max_row,
            x,
            y,
            w,
            h,
            round(area, 2),
        ])

    # ヘッダー装飾
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="2563EB",
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col in "ABCDEF":
        ws.column_dimensions[col].width = 15

    # ファイルに書き込まずメモリ上でExcel生成
    output = BytesIO()
    wb.save(output)

    return output.getvalue()
